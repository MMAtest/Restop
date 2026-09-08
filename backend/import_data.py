"""
Import Mercuriale + Recettes - La Table d'Augustine
Matching exact par nom (case-insensitive). Creation si introuvable.
"""
import asyncio
import os
import uuid
import re
from datetime import datetime
import openpyxl
import motor.motor_asyncio

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")
client = motor.motor_asyncio.AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]

CATEGORY_MAP = {
    "Boucherie": "boucherie",
    "Fromagerie": "fromagerie",
    "Marée": "marée",
    "Extra": "extra",
}


async def find_product_by_name(name, all_produits, all_preparations):
    """Cherche un produit ou preparation par nom exact (case-insensitive)"""
    name_lower = name.lower().strip()
    
    # 1. Match exact dans produits
    for p in all_produits:
        if p["nom"].lower().strip() == name_lower:
            return p, "produit"
    
    # 2. Match exact dans preparations
    for p in all_preparations:
        if p["nom"].lower().strip() == name_lower:
            return p, "preparation"
    
    return None, None


async def import_mercuriale(filepath):
    print("\n========== IMPORT MERCURIALE ==========")
    wb = openpyxl.load_workbook(filepath)
    
    stats = {"produits_crees": 0, "produits_maj": 0, "doublons": 0, "fournisseurs_crees": 0}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        
        headers = [str(h).strip().lower() if h else '' for h in rows[0]]
        data_rows = [r for r in rows[1:] if r and r[0]]
        categorie = CATEGORY_MAP.get(sheet_name, "extra")
        
        print(f"\n--- {sheet_name} -> {categorie} ({len(data_rows)} lignes) ---")
        
        nom_idx = next((i for i, h in enumerate(headers) if 'nom' in h), 0)
        prix_idx = next((i for i, h in enumerate(headers) if 'prix' in h), 1)
        unite_idx = next((i for i, h in enumerate(headers) if 'unit' in h), 2)
        fourn_idx = next((i for i, h in enumerate(headers) if 'fourn' in h), 3)
        
        for row in data_rows:
            nom = str(row[nom_idx]).strip()
            if not nom:
                continue
            
            try:
                prix = float(row[prix_idx]) if row[prix_idx] and str(row[prix_idx]).replace('.','').replace(',','').isdigit() else 0
            except:
                prix = 0
            
            unite_raw = str(row[unite_idx]).strip() if len(row) > unite_idx and row[unite_idx] else "kg"
            unite_map = {"kg": "kg", "l": "L", "piece": "piece", "pièce": "piece", "baguette": "piece"}
            unite = unite_map.get(unite_raw.lower(), unite_raw)
            
            fourn_nom = str(row[fourn_idx]).strip() if len(row) > fourn_idx and row[fourn_idx] else ""
            
            # Creer fournisseur si besoin
            if fourn_nom and fourn_nom != 'None':
                existing_f = await db.fournisseurs.find_one({"nom": re.compile(f"^{re.escape(fourn_nom)}$", re.IGNORECASE)})
                if not existing_f:
                    await db.fournisseurs.insert_one({
                        "id": str(uuid.uuid4()), "nom": fourn_nom,
                        "categories": [categorie], "contact": "", "telephone": "",
                        "email": "", "heure_limite_commande": "", "delai_livraison": "",
                        "minimum_commande": "", "notes": "",
                        "created_at": datetime.utcnow()
                    })
                    stats["fournisseurs_crees"] += 1
                    print(f"  + Fournisseur: {fourn_nom}")
                else:
                    cats = existing_f.get("categories", [])
                    if categorie not in cats:
                        cats.append(categorie)
                        await db.fournisseurs.update_one({"id": existing_f["id"]}, {"$set": {"categories": cats}})
            
            # Chercher produit existant par nom exact
            existing_p = await db.produits.find_one({"nom": re.compile(f"^{re.escape(nom)}$", re.IGNORECASE)})
            
            if existing_p:
                old_prix = existing_p.get("prix_achat", 0) or 0
                if prix > 0 and (old_prix == 0 or prix < old_prix):
                    await db.produits.update_one({"id": existing_p["id"]}, {"$set": {
                        "prix_achat": prix, "fournisseur": fourn_nom or existing_p.get("fournisseur", ""),
                        "categorie": categorie
                    }})
                    stats["produits_maj"] += 1
                else:
                    stats["doublons"] += 1
            else:
                produit_id = str(uuid.uuid4())
                await db.produits.insert_one({
                    "id": produit_id, "nom": nom, "categorie": categorie,
                    "unite": unite, "prix_achat": prix, "fournisseur": fourn_nom,
                    "stock_min": 0, "stock_actuel": 0, "dlc": None,
                    "notes": "", "archived": False, "created_at": datetime.utcnow()
                })
                await db.stock_ingredients.insert_one({
                    "id": str(uuid.uuid4()), "produit_id": produit_id,
                    "produit_nom": nom, "quantite": 0, "unite": unite,
                    "prix_unitaire": prix, "dlc": None, "emplacement": "reserve",
                    "updated_at": datetime.utcnow()
                })
                stats["produits_crees"] += 1
    
    print(f"\n--- RESULTAT MERCURIALE ---")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    return stats


async def import_recettes(filepath):
    print("\n\n========== IMPORT RECETTES ==========")
    wb = openpyxl.load_workbook(filepath)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    
    headers = [str(h).strip() if h else '' for h in rows[0]]
    data_rows = rows[1:]
    
    col = {}
    for i, h in enumerate(headers):
        hl = h.lower()
        if 'nom' in hl and 'recette' in hl: col['nom'] = i
        elif 'categorie' in hl or 'catégorie' in hl: col['cat'] = i
        elif 'portion' in hl: col['portions'] = i
        elif 'prix' in hl: col['prix'] = i
        elif 'temps' in hl: col['temps'] = i
        elif 'description' in hl: col['desc'] = i
        elif 'ingredient' in hl or 'ingrédient' in hl: col['ing'] = i
        elif 'quantite' in hl or 'quantité' in hl: col['qty'] = i
        elif 'unite' in hl or 'unité' in hl: col['unit'] = i
    
    print(f"Colonnes: {col}")
    
    # Charger tous les produits et preparations
    all_produits = await db.produits.find({}, {"_id": 0}).to_list(length=5000)
    all_preparations = await db.preparations.find({}, {"_id": 0}).to_list(length=5000)
    print(f"Produits: {len(all_produits)}, Preparations: {len(all_preparations)}")
    
    # Regrouper par recette
    recettes = {}
    for row in data_rows:
        if not row or not row[col['nom']]:
            continue
        
        nom = str(row[col['nom']]).strip()
        if nom not in recettes:
            recettes[nom] = {
                "nom": nom,
                "categorie": str(row[col.get('cat', 1)] or "").strip() or None,
                "portions": int(row[col.get('portions', 2)] or 1),
                "prix_vente": float(row[col.get('prix', 3)] or 0),
                "temps_preparation": int(row[col.get('temps', 4)] or 0) if row[col.get('temps', 4)] else None,
                "description": str(row[col.get('desc', 5)] or "").strip() or None,
                "ingredients": []
            }
        
        ing_nom = str(row[col.get('ing', 6)] or "").strip()
        qty = float(row[col.get('qty', 7)] or 0)
        unite = str(row[col.get('unit', 8)] or "").strip()
        
        if not ing_nom or qty <= 0:
            continue
        
        # Chercher par nom exact
        match, match_type = await find_product_by_name(ing_nom, all_produits, all_preparations)
        
        if match:
            ing_data = {
                "produit_id": match["id"],
                "produit_nom": match["nom"],
                "ingredient_id": match["id"],
                "ingredient_type": match_type,
                "ingredient_nom": match["nom"],
                "quantite": qty,
                "unite": unite
            }
        else:
            # Creer le produit
            new_id = str(uuid.uuid4())
            new_prod = {
                "id": new_id, "nom": ing_nom, "categorie": "extra",
                "unite": unite if unite in ["kg", "L", "piece"] else "kg",
                "prix_achat": 0, "fournisseur": "", "stock_min": 0,
                "stock_actuel": 0, "dlc": None, "notes": "",
                "archived": False, "created_at": datetime.utcnow()
            }
            await db.produits.insert_one(new_prod)
            await db.stock_ingredients.insert_one({
                "id": str(uuid.uuid4()), "produit_id": new_id,
                "produit_nom": ing_nom, "quantite": 0,
                "unite": new_prod["unite"], "prix_unitaire": 0,
                "dlc": None, "emplacement": "reserve",
                "updated_at": datetime.utcnow()
            })
            all_produits.append(new_prod)
            print(f"  + Produit cree: '{ing_nom}'")
            
            ing_data = {
                "produit_id": new_id, "produit_nom": ing_nom,
                "ingredient_id": new_id, "ingredient_type": "produit",
                "ingredient_nom": ing_nom, "quantite": qty, "unite": unite
            }
        
        recettes[nom]["ingredients"].append(ing_data)
    
    # Sauvegarder les recettes
    created = 0
    for nom, data in recettes.items():
        # Calculer cout matiere
        cout = 0
        for ing in data["ingredients"]:
            p = await db.produits.find_one({"id": ing.get("produit_id", "")}, {"_id": 0})
            if p and p.get("prix_achat"):
                cout += ing["quantite"] * p["prix_achat"]
        
        data["cout_matiere"] = round(cout, 2)
        portions = data.get("portions", 1) or 1
        if data.get("prix_vente") and cout > 0:
            data["coefficient_reel"] = round((cout / portions) / data["prix_vente"], 2)
        
        existing = await db.recettes.find_one({"nom": nom})
        if existing:
            await db.recettes.update_one({"nom": nom}, {"$set": data})
        else:
            data["id"] = str(uuid.uuid4())
            data["instructions"] = None
            data["coefficient_prevu"] = None
            data["created_at"] = datetime.utcnow()
            await db.recettes.insert_one(data)
            created += 1
    
    print(f"\n--- RESULTAT RECETTES ---")
    print(f"  Recettes creees: {created}")
    print(f"  Total: {len(recettes)}")
    return created


async def main():
    print("=" * 50)
    print("IMPORT La Table d'Augustine")
    print("=" * 50)
    
    if os.path.exists("/tmp/mercuriale.xlsx"):
        await import_mercuriale("/tmp/mercuriale.xlsx")
    
    if os.path.exists("/tmp/menu.xlsx"):
        await import_recettes("/tmp/menu.xlsx")
    
    # Stats finales
    total_p = await db.produits.count_documents({})
    total_f = await db.fournisseurs.count_documents({})
    total_r = await db.recettes.count_documents({})
    print(f"\n=== ETAT FINAL ===")
    print(f"  Produits: {total_p}")
    print(f"  Fournisseurs: {total_f}")
    print(f"  Recettes: {total_r}")

if __name__ == "__main__":
    asyncio.run(main())
